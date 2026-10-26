"""s4 渲染：部门报告 + 各组报告 + 系数表 + 收件人清单。"""
from __future__ import annotations

import csv
import json
import os

from lib.render_dept import render as render_dept
from lib.render_group import PALETTE, render as render_group


def _recipients(A, cfg, path):
    """各组收件人清单：owner 附带部门报告，成员只带本组报告。"""
    groups = {g["name"]: g for g in A["groups"]}
    rows = []
    for g in A["groups"]:
        owner = g["name"]
        members = sorted([p for p in A["people"] if p["owner"] == owner],
                         key=lambda x: -x["credits"])
        for p in members:
            if not p["email"]:
                continue
            rows.append({
                "group": owner,
                "owner": owner,
                "name": p["name"],
                "email": p["email"],
                "role": "owner" if p["name"] == owner else "member",
                "report": f"team-{owner}.html",
                "dept_report": "token_credits_report.html",
                "send_dept": "true" if p["name"] == owner else "false",
            })
        if not any(r["name"] == owner and r["group"] == owner for r in rows):
            # owner 自己不在组里（或没邮箱），补一行
            rows.append({"group": owner, "owner": owner, "name": owner, "email": "",
                         "role": "owner", "report": f"team-{owner}.html",
                         "dept_report": "token_credits_report.html", "send_dept": "true"})
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["group", "owner", "name", "email", "role",
                                          "report", "dept_report", "send_dept"])
        w.writeheader()
        w.writerows(rows)
    return len(rows)


def _coefficient_md(A, cfg, path):
    lines = [f"# Token ↔ Credits 系数 · {A['meta']['month']}", "",
             f"数据区间 {A['meta']['data_start']} → {A['meta']['data_end']}，"
             f"明细 {A['meta']['rows']:,} 行。单位：Credits / 1,000,000 tokens。", "",
             "| 模型 | 时段 | 档位 | token in | token out | cache | cache read | cache write |",
             "|---|---|---|---|---|---|---|---|"]
    combos = {}
    for r in A["rates"]:
        combos.setdefault((r["model"], r["ctx"], r["period"]), {})[r["role"]] = r["rate"]

    def fnum(v):
        return "—" if v is None else (f"{v:,.0f}" if abs(v - round(v)) < 1e-6 else f"{v:,.2f}")

    for (m, ctx, per), cells in combos.items():
        lines.append(f"| {m} | {per} | {ctx or '—'} | {fnum(cells.get('token_in'))} | "
                     f"{fnum(cells.get('token_out'))} | {fnum(cells.get('cache'))} | "
                     f"{fnum(cells.get('cache_read'))} | {fnum(cells.get('cache_write_5m'))} |")
    open(path, "w", encoding="utf-8").write("\n".join(lines) + "\n")


def _coefficient_xlsx(A, cfg, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "系数总表"
    ws.append(["模型", "时段", "档位", "token in", "token out", "cache", "cache read", "cache write(5m)", "样本行数"])
    combos = {}
    for r in A["rates"]:
        d = combos.setdefault((r["model"], r["ctx"], r["period"]), {"cells": {}, "rows": 0})
        d["cells"][r["role"]] = r["rate"]
        d["rows"] += r["rows"]

    def rnd(v):
        return None if v is None else (round(v) if abs(v - round(v)) < 1e-6 else round(v, 4))

    for (m, ctx, per), d in combos.items():
        c = d["cells"]
        ws.append([m, per, ctx or "—", rnd(c.get("token_in")), rnd(c.get("token_out")),
                   rnd(c.get("cache")), rnd(c.get("cache_read")), rnd(c.get("cache_write_5m")), d["rows"]])
    hdr = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr
    for i, w in enumerate([26, 10, 12, 11, 11, 10, 11, 15, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("逐人")
    ws2.append(["姓名", "共享人", "组", "部门", "邮箱", "席位", "额度", "本月点数", "席位消耗",
                "加油包", "结余点数", "折合金额", "下月预测", "状态"])
    for p in sorted(A["people"], key=lambda x: -x["credits"]):
        ws2.append([p["name"], p["shared_with"], p["owner"], p["dept"], p["email"],
                    "高级" if p["premium"] else "标准", p["quota"], round(p["credits"], 2),
                    round(p["seat"], 2), round(p["pack"], 2), round(p["unused"], 2),
                    round(p["waste"], 2), round(p["forecast"], 2), p["status"]])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr
    for i, w in enumerate([14, 14, 10, 10, 30, 8, 10, 12, 12, 12, 12, 10, 12, 12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("模型")
    ws3.append(["模型", "点数", "占比", "tokens", "混合单价/1M", "覆盖人数", "覆盖组数"])
    total = A["meta"]["total_credits"] or 1
    for m in A["models"]:
        ws3.append([m["name"], round(m["credits"], 2), m["credits"] / total,
                    round(m["tokens"]), round(m["mix_price"], 2), m["users"], len(m["groups"])])
    for cell in ws3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr
    for i, w in enumerate([26, 14, 8, 16, 14, 10, 10], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    return True


def run(cfg, A, out_dir, log):
    log("─" * 68)
    log("s4 渲染报告")
    files = []
    total = A["meta"]["total_credits"] or 1

    dept = os.path.join(out_dir, "token_credits_report.html")
    open(dept, "w", encoding="utf-8").write(render_dept(A, cfg))
    files.append(dept)
    log(f"  ✓ {os.path.basename(dept)}  {os.path.getsize(dept) // 1024} KB")

    if cfg.output.get("group_reports", True):
        for i, g in enumerate(A["groups"]):
            p = os.path.join(out_dir, f"team-{g['name']}.html")
            open(p, "w", encoding="utf-8").write(render_group(A, cfg, g["name"], PALETTE[i % len(PALETTE)]))
            files.append(p)
            log(f"  ✓ {os.path.basename(p):<24} {g['count']:>2} 席  {g['credits']:>10,.0f} 点"
                f"  {g['credits'] / total * 100:>5.1f}%")

    rec = os.path.join(out_dir, "recipients.csv")
    n = _recipients(A, cfg, rec)
    files.append(rec)
    log(f"  ✓ recipients.csv  {n} 行")

    md = os.path.join(out_dir, "token_credits_coefficients.md")
    _coefficient_md(A, cfg, md)
    files.append(md)

    xlsx = os.path.join(out_dir, "token_credits_coefficients.xlsx")
    if _coefficient_xlsx(A, cfg, xlsx):
        files.append(xlsx)
        log(f"  ✓ token_credits_coefficients.xlsx")
    else:
        log("  · 跳过 xlsx（未安装 openpyxl）")

    return files