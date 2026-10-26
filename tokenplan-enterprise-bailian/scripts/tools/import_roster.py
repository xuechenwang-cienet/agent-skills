#!/usr/bin/env python3
"""Import an xlsx roster into config/members.csv (one-off / refresh helper).

Usage:
    python3 scripts/tools/import_roster.py --xlsx "../../202610_usage.xlsx"
    python3 scripts/tools/import_roster.py --xlsx ... --accounts accounts.json

`accounts.json` is an optional list of {"id": "acc_xxx", "name": "<Bailian 显示名>"}
used to fill the `bailian_accounts` column. Accounts are matched to roster rows by
normalising these fields: Bailian(old) / DingTalk name / Chinese name / English name.
Matching is only used *here* (import time); the runtime never guesses.
"""
import argparse, csv, json, os, re, sys

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def norm(s):
    return re.sub(r"[\s_\-\.]", "", str(s or "")).lower()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="roster workbook (.xlsx)")
    ap.add_argument("--sheet", default=None, help="sheet name (default: first)")
    ap.add_argument("--accounts", default=None, help="accounts.json with Bailian account ids/names")
    ap.add_argument("--out", default=os.path.join(BASE, "config", "members.csv"))
    a = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        sys.exit("需要 openpyxl 才能导入 xlsx：uv pip install --python ~/.venvs/tools/bin/python openpyxl")

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]

    def col(row, *names):
        for n in names:
            for k in hdr:
                if k and str(k).strip().lower() == n.lower():
                    return row.get(k)
        return None

    out = []
    for r in rows:
        seat = col(r, "Seat token") or 25000
        try:
            seat = int(float(seat))
        except Exception:
            seat = 25000
        name_zh = (col(r, "Chinese Name") or "").strip()
        ding = (col(r, "DingTalk Name") or "").strip()
        name_en = (col(r, "Name") or "").strip()
        out.append(dict(
            hrid=str(col(r, "HRID") or "").strip(),
            name_en=name_en,
            name_zh=name_zh or ding or name_en,
            dept=(col(r, "Dept.") or "").strip() or "—",
            pcode=(col(r, "P-code") or "").strip() or "—",
            email=(col(r, "CIeNET Email") or "").strip(),
            team_owner=(col(r, "Team Owner") or "").strip(),
            seat_type="premium" if seat > 25000 else "standard",
            seat_credits=seat,
            sharing_with=(col(r, "Sharing with") or "").strip(),
            bailian_accounts="",
            bailian_account_ids="",
            active="true",
            _keys=[col(r, "Name in Bailian(Old)"), col(r, "DingTalk Name"),
                   col(r, "Chinese Name"), col(r, "Name")],
        ))

    # ---- attach Bailian accounts ----
    if a.accounts:
        accs = json.load(open(a.accounts, encoding="utf-8"))
        if isinstance(accs, dict):
            accs = [{"id": k, "name": v} for k, v in accs.items()]
        idx = {}
        for row in out:
            for k in row["_keys"]:
                if k:
                    idx.setdefault(norm(k), row)
        unmatched = []
        for acc in accs:
            row = idx.get(norm(acc["name"]))
            if not row:
                unmatched.append(acc)
                continue
            cur = [x for x in row["bailian_accounts"].split(";") if x]
            if acc["name"] not in cur:
                cur.append(acc["name"])
            row["bailian_accounts"] = ";".join(cur)
            ids = [x for x in row["bailian_account_ids"].split(";") if x]
            if acc["id"] not in ids:
                ids.append(acc["id"])
            row["bailian_account_ids"] = ";".join(ids)
        if unmatched:
            print("⚠ 未能匹配到名册的账号：", file=sys.stderr)
            for u in unmatched:
                print(f"   {u['id']}  {u['name']}", file=sys.stderr)

    cols = ["hrid", "name_en", "name_zh", "dept", "pcode", "email", "team_owner",
            "seat_type", "seat_credits", "sharing_with", "bailian_accounts",
            "bailian_account_ids", "active"]
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    with open(a.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for row in out:
            w.writerow({k: row[k] for k in cols})

    print(f"✓ 写入 {a.out}：{len(out)} 行")
    print(f"  有 Bailian 账号：{sum(1 for r in out if r['bailian_accounts'])}")
    print(f"  无账号（未开通）：{sum(1 for r in out if not r['bailian_accounts'])}")
    print(f"  共享席位：{sum(1 for r in out if r['sharing_with'])}")
    print(f"  项目组：{len({r['team_owner'] for r in out if r['team_owner']})}")


if __name__ == "__main__":
    main()